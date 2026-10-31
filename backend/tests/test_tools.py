import io
import json
import pytest
import tempfile
import os
from pathlib import Path
from app.services.tools import registry
from app.core.config import settings

@pytest.mark.asyncio
async def test_registry_contains_12_tools():
    tools = registry.list_tools()
    assert len(tools) == 12
    tool_ids = {t.id for t in tools}
    expected = {
        "pdf_extractor", "docx_reader", "excel_processor", "csv_processor",
        "json_parser", "calculator", "datetime_tool", "chart_generator",
        "url_validator", "data_cleaner", "duplicate_finder", "text_summarizer"
    }
    assert tool_ids == expected

def test_gemini_declarations_schema_compliance():
    """Verify Gemini function declarations schema rules: all ARRAY types must have items."""
    declarations = registry.to_gemini_declarations()
    assert len(declarations) == 12

    for decl in declarations:
        assert "name" in decl
        assert "description" in decl
        assert "parameters" in decl
        params = decl["parameters"]
        assert params["type"] == "OBJECT"
        props = params.get("properties", {})
        for p_name, p_val in props.items():
            assert "type" in p_val
            assert p_val["type"] in {"STRING", "NUMBER", "INTEGER", "BOOLEAN", "ARRAY", "OBJECT"}
            if p_val["type"] == "ARRAY":
                assert "items" in p_val, f"ARRAY property '{decl['name']}.{p_name}' must have 'items'"
                assert isinstance(p_val["items"], dict)
                assert "type" in p_val["items"]
                assert p_val["items"]["type"] in {"STRING", "NUMBER", "INTEGER", "BOOLEAN", "ARRAY", "OBJECT"}

@pytest.mark.asyncio
async def test_calculator_tool_operations():
    calc = registry.get("calculator")
    assert calc is not None

    # Addition
    res = await calc.execute({"expression": "125 + 875"})
    assert res.success is True
    assert res.data["result"] == 1000

    # Power & Math functions
    res = await calc.execute({"expression": "sqrt(144) + 2**4"})
    assert res.success is True
    assert res.data["result"] == 28

    # Percentage
    res = await calc.execute({"expression": "25% of 200"})
    assert res.success is True
    assert res.data["result"] == 50

    # Division by zero
    res = await calc.execute({"expression": "100 / 0"})
    assert res.success is False
    assert "Division by zero" in (res.error or "")

    # Security: disallowed AST operations
    res = await calc.execute({"expression": "__import__('os').system('ls')"})
    assert res.success is False

@pytest.mark.asyncio
async def test_datetime_tool_operations():
    dt_tool = registry.get("datetime_tool")
    assert dt_tool is not None

    # Current time
    res = await dt_tool.execute({"operation": "now", "timezone": "UTC"})
    assert res.success is True
    assert "iso" in res.data

    # Date difference
    res = await dt_tool.execute({
        "operation": "diff",
        "date1": "2026-05-01",
        "date2": "2026-05-15"
    })
    assert res.success is True
    assert res.data["days_difference"] == 14

    # Add/subtract
    res = await dt_tool.execute({
        "operation": "add_subtract",
        "date1": "2026-01-01",
        "days": 10
    })
    assert res.success is True
    assert res.data["result_date"] == "2026-01-11"

@pytest.mark.asyncio
async def test_url_validator_and_ssrf_protection():
    validator = registry.get("url_validator")
    assert validator is not None

    # Syntax valid public URL
    res = await validator.execute({"url": "https://example.com/api/test", "check_accessibility": False})
    assert res.success is True
    assert res.data["is_valid_syntax"] is True

    # Invalid URL syntax
    res = await validator.execute({"url": "htp:/invalid-url", "check_accessibility": False})
    assert res.success is False

    # SSRF Protection: Loopback / Private IP should be blocked
    res = await validator.execute({"url": "http://127.0.0.1:8080/admin", "check_accessibility": True})
    assert res.success is False
    assert any(term in (res.error or "") for term in ["localhost", "loopback", "private", "forbidden", "blocked"])

    res = await validator.execute({"url": "http://192.168.1.1/setup", "check_accessibility": True})
    assert res.success is False
    assert any(term in (res.error or "") for term in ["private", "RFC-1918", "forbidden", "blocked"])

@pytest.mark.asyncio
async def test_json_parser_tool():
    parser = registry.get("json_parser")
    assert parser is not None

    valid_json = json.dumps({
        "agent": {
            "name": "Precision Assistant",
            "version": 2,
            "tools": ["pdf_extractor", "calculator"]
        }
    })

    # Validate and parse
    res = await parser.execute({"json_content": valid_json})
    assert res.success is True
    assert res.data["valid"] is True

    # Nested path extraction
    res = await parser.execute({"json_content": valid_json, "extract_path": "agent.name"})
    assert res.success is True
    assert res.data["extracted_value"] == "Precision Assistant"

    # Syntax error
    res = await parser.execute({"json_content": "{invalid json: true"})
    assert res.success is False
    assert "JSON Syntax Error" in (res.error or "")

@pytest.mark.asyncio
async def test_csv_processor_tool():
    csv_tool = registry.get("csv_processor")
    assert csv_tool is not None

    raw_csv = "Name,Department,Salary\nAlice,Engineering,120000\nBob,Marketing,90000\nCharlie,Engineering,110000\n"
    res = await csv_tool.execute({
        "csv_content": raw_csv,
        "operation": "filter",
        "filter_column": "Department",
        "filter_operator": "equals",
        "filter_value": "Engineering"
    })
    assert res.success is True
    assert res.data["row_count"] == 2
    assert "Alice" in str(res.data["sample_data"])
    assert "Bob" not in str(res.data["sample_data"])

@pytest.mark.asyncio
async def test_data_cleaner_and_duplicate_finder():
    cleaner = registry.get("data_cleaner")
    finder = registry.get("duplicate_finder")

    dirty_csv = " Full Name , Age , City \n John Doe , 30 , New York \n John Doe , 30 , New York \nJane Smith, 25, London\n"
    
    # Cleaner
    res_clean = await cleaner.execute({
        "csv_content": dirty_csv,
        "remove_duplicates": True,
        "standardize_column_names": True,
        "normalize_whitespace": True
    })
    assert res_clean.success is True
    assert res_clean.data["original_row_count"] == 3
    assert res_clean.data["cleaned_row_count"] == 2
    assert "full_name" in res_clean.data["columns"]

    # Duplicate finder
    res_dup = await finder.execute({"csv_content": dirty_csv})
    assert res_dup.success is True
    assert res_dup.data["duplicate_row_count"] == 1
    assert res_dup.data["unique_record_count"] == 2

@pytest.mark.asyncio
async def test_chart_generator_tool():
    chart_tool = registry.get("chart_generator")
    assert chart_tool is not None

    data = [
        {"Quarter": "Q1", "Revenue": 15000},
        {"Quarter": "Q2", "Revenue": 23000},
        {"Quarter": "Q3", "Revenue": 19000},
        {"Quarter": "Q4", "Revenue": 32000}
    ]

    res = await chart_tool.execute({
        "data": data,
        "chart_type": "bar",
        "x_field": "Quarter",
        "y_field": "Revenue",
        "title": "Quarterly Revenue"
    })
    assert res.success is True
    assert res.download_url is not None
    assert res.data["base64_image"].startswith("data:image/png;base64,")

@pytest.mark.asyncio
async def test_pdf_and_docx_tools():
    pdf_tool = registry.get("pdf_extractor")
    docx_tool = registry.get("docx_reader")

    # Missing file handling
    res = await pdf_tool.execute({"file_path": "non_existent_file.pdf"})
    assert res.success is False

    os.makedirs(settings.STORAGE_DIR, exist_ok=True)

    # Generate synthetic PDF with PyMuPDF in STORAGE_DIR
    import pymupdf
    pdf_file = os.path.join(settings.STORAGE_DIR, "test_synthetic.pdf")
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Hello AGENT.STUDIO PDF Extractor!")
    doc.save(pdf_file)
    doc.close()

    try:
        res_pdf = await pdf_tool.execute({"file_path": pdf_file})
        assert res_pdf.success is True
        assert res_pdf.data["total_pages"] == 1
        assert "Hello AGENT.STUDIO PDF Extractor!" in res_pdf.data["pages"][0]["text"]
    finally:
        if os.path.exists(pdf_file):
            os.remove(pdf_file)

    # Generate synthetic DOCX with python-docx in STORAGE_DIR
    import docx
    docx_file = os.path.join(settings.STORAGE_DIR, "test_synthetic.docx")
    doc_word = docx.Document()
    doc_word.add_heading("Agent Studio Specification", level=1)
    doc_word.add_paragraph("This is a test paragraph for DOCX reader.")
    doc_word.save(docx_file)

    try:
        res_docx = await docx_tool.execute({"file_path": docx_file})
        assert res_docx.success is True
        assert len(res_docx.data["headings"]) >= 1
        assert "Agent Studio Specification" in res_docx.data["headings"][0]["text"]
        assert "This is a test paragraph for DOCX reader." in res_docx.data["paragraphs"][0]
    finally:
        if os.path.exists(docx_file):
            os.remove(docx_file)

@pytest.mark.asyncio
async def test_excel_processor_tool():
    excel_tool = registry.get("excel_processor")
    assert excel_tool is not None

    os.makedirs(settings.STORAGE_DIR, exist_ok=True)
    import pandas as pd
    excel_file = os.path.join(settings.STORAGE_DIR, "test_metrics.xlsx")
    df = pd.DataFrame({
        "Product": ["Alpha", "Beta", "Gamma"],
        "Sales": [120, 250, 180]
    })
    df.to_excel(excel_file, index=False, sheet_name="Overview")

    try:
        res = await excel_tool.execute({
            "file_path": excel_file,
            "sheet_name": "Overview",
            "operation": "info"
        })
        assert res.success is True
        assert res.data["sheet_names"] == ["Overview"]
        assert "Product" in res.data["columns"]
        assert res.data["total_rows"] == 3
    finally:
        if os.path.exists(excel_file):
            os.remove(excel_file)

# ====================================================================
# API Endpoint Tests
# ====================================================================

def test_api_list_tools(client, admin_headers):
    resp = client.get("/api/tools", headers=admin_headers)
    assert resp.status_code == 200
    data = resp.json()
    assert len(data) == 12
    categories = {t["category"] for t in data}
    expected_categories = {
        "Documents & Files", "Data Processing", "Data Utilities",
        "Validation & Utilities", "Gemini AI Tools"
    }
    assert categories == expected_categories

def test_api_toggle_tool_permission(client, admin_headers, creator_headers):
    # Non-admin cannot toggle
    resp = client.patch("/api/tools/calculator/toggle", json={"is_enabled": False}, headers=creator_headers)
    assert resp.status_code == 403

    # Admin can toggle
    resp_admin = client.patch("/api/tools/calculator/toggle", json={"is_enabled": False}, headers=admin_headers)
    assert resp_admin.status_code == 200
    assert resp_admin.json()["is_enabled"] is False

    # Toggle back
    resp_back = client.patch("/api/tools/calculator/toggle", json={"is_enabled": True}, headers=admin_headers)
    assert resp_back.status_code == 200
    assert resp_back.json()["is_enabled"] is True

def test_api_test_tool_endpoint(client, admin_headers):
    # Interactive test of calculator
    resp = client.post("/api/tools/calculator/test", json={
        "tool_id": "calculator",
        "parameters": {"expression": "50 * 4"}
    }, headers=admin_headers)
    assert resp.status_code == 200
    body = resp.json()
    assert body["success"] is True
    assert body["data"]["result"] == 200

def test_api_download_nonexistent_and_traversal(client, admin_headers):
    # Non-existent file should be 404
    resp = client.get("/api/tools/download/nonexistent_file.png", headers=admin_headers)
    assert resp.status_code == 404

    # Traversal in name should be 400
    resp_bad = client.get("/api/tools/download/..bad_file.png", headers=admin_headers)
    assert resp_bad.status_code == 400
