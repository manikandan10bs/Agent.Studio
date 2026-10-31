import io
import os
import csv
import json
from typing import Dict, Any, Optional, List
import pandas as pd
from app.services.tools.registry import BaseTool, ToolExecutionResult
from app.services.tools.utils import resolve_safe_file_path, save_output_file

# ====================================================================
# 1. PDF Extractor
# ====================================================================
class PdfExtractorTool(BaseTool):
    id = "pdf_extractor"
    name = "PDF Extractor"
    description = "Extracts text page-by-page from uploaded PDF documents, preserving page numbers and metadata. Handles encrypted or invalid files with clear errors."
    category = "Documents & Files"
    icon = "FileText"
    timeout_seconds = 30
    permission_required = "creator"

    input_schema = {
        "type": "object",
        "properties": {
            "file_path": {
                "type": "string",
                "description": "The file name or path of the uploaded PDF to extract text from."
            },
            "start_page": {
                "type": "integer",
                "description": "Optional 1-indexed page number to start extracting from."
            },
            "end_page": {
                "type": "integer",
                "description": "Optional 1-indexed page number to stop extracting."
            }
        },
        "required": ["file_path"]
    }

    output_schema = {
        "type": "object",
        "properties": {
            "total_pages": {"type": "integer"},
            "extracted_pages": {"type": "integer"},
            "pages": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "page": {"type": "integer"},
                        "text": {"type": "string"},
                        "char_count": {"type": "integer"}
                    }
                }
            },
            "metadata": {"type": "object"}
        }
    }

    async def execute(self, params: Dict[str, Any], context: Optional[Dict[str, Any]] = None) -> ToolExecutionResult:
        file_path = params.get("file_path")
        if not file_path:
            return ToolExecutionResult(success=False, tool_id=self.id, error="Missing required parameter 'file_path'.")

        try:
            abs_path = resolve_safe_file_path(file_path)
        except Exception as e:
            return ToolExecutionResult(success=False, tool_id=self.id, error=str(e))

        start_page = params.get("start_page") or 1
        end_page = params.get("end_page")

        pages_data = []
        metadata = {}
        total_pages = 0

        # Try PyMuPDF first
        try:
            import fitz
            doc = fitz.open(abs_path)
            if doc.is_encrypted:
                return ToolExecutionResult(
                    success=False,
                    tool_id=self.id,
                    error=f"PDF '{os.path.basename(abs_path)}' is password-protected or encrypted. Please provide an unencrypted PDF."
                )

            total_pages = len(doc)
            if total_pages == 0:
                return ToolExecutionResult(
                    success=False,
                    tool_id=self.id,
                    error=f"PDF document is empty (0 pages)."
                )

            metadata = {
                "format": doc.metadata.get("format", "PDF"),
                "title": doc.metadata.get("title", ""),
                "author": doc.metadata.get("author", ""),
                "page_count": total_pages
            }

            max_p = min(end_page or total_pages, total_pages)
            min_p = max(1, start_page)

            for p_num in range(min_p, max_p + 1):
                page = doc.load_page(p_num - 1)
                text = page.get_text("text").strip()
                pages_data.append({
                    "page": p_num,
                    "text": text,
                    "char_count": len(text)
                })
            doc.close()

        except Exception as fitz_err:
            # Fallback to pypdf
            try:
                from pypdf import PdfReader
                reader = PdfReader(abs_path)
                if reader.is_encrypted:
                    return ToolExecutionResult(
                        success=False,
                        tool_id=self.id,
                        error=f"PDF '{os.path.basename(abs_path)}' is encrypted. Please provide an unencrypted document."
                    )
                total_pages = len(reader.pages)
                metadata = {"page_count": total_pages}
                max_p = min(end_page or total_pages, total_pages)
                min_p = max(1, start_page)

                for p_num in range(min_p, max_p + 1):
                    page = reader.pages[p_num - 1]
                    text = (page.extract_text() or "").strip()
                    pages_data.append({
                        "page": p_num,
                        "text": text,
                        "char_count": len(text)
                    })
            except Exception as pypdf_err:
                return ToolExecutionResult(
                    success=False,
                    tool_id=self.id,
                    error=f"Failed to read PDF file '{os.path.basename(abs_path)}': {str(pypdf_err)}"
                )

        preview = f"Extracted {len(pages_data)} page(s) from '{os.path.basename(abs_path)}'. Total pages: {total_pages}."
        return ToolExecutionResult(
            success=True,
            tool_id=self.id,
            data={
                "total_pages": total_pages,
                "extracted_pages": len(pages_data),
                "pages": pages_data,
                "metadata": metadata
            },
            output_preview=preview
        )


# ====================================================================
# 2. Word Document Reader
# ====================================================================
class DocxReaderTool(BaseTool):
    id = "docx_reader"
    name = "Word Document Reader"
    description = "Reads Microsoft Word (.docx) documents, extracting structured paragraphs, headings by level, and tables formatted as clean rows."
    category = "Documents & Files"
    icon = "FileCode"
    timeout_seconds = 25
    permission_required = "creator"

    input_schema = {
        "type": "object",
        "properties": {
            "file_path": {
                "type": "string",
                "description": "File name or path of the uploaded .docx document."
            },
            "include_tables": {
                "type": "boolean",
                "description": "Whether to extract structured tables (default true)."
            }
        },
        "required": ["file_path"]
    }

    output_schema = {
        "type": "object",
        "properties": {
            "total_paragraphs": {"type": "integer"},
            "total_headings": {"type": "integer"},
            "total_tables": {"type": "integer"},
            "headings": {"type": "array"},
            "paragraphs": {"type": "array"},
            "tables": {"type": "array"}
        }
    }

    async def execute(self, params: Dict[str, Any], context: Optional[Dict[str, Any]] = None) -> ToolExecutionResult:
        file_path = params.get("file_path")
        if not file_path:
            return ToolExecutionResult(success=False, tool_id=self.id, error="Missing required parameter 'file_path'.")

        try:
            abs_path = resolve_safe_file_path(file_path)
            from docx import Document as DocxDocument
            doc = DocxDocument(abs_path)
        except Exception as e:
            return ToolExecutionResult(success=False, tool_id=self.id, error=f"Could not open DOCX document: {str(e)}")

        include_tables = params.get("include_tables", True)

        headings = []
        paragraphs = []
        tables_data = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            style_name = p.style.name if p.style else ""
            if "Heading" in style_name:
                level = 1
                try:
                    level = int(style_name.replace("Heading", "").strip())
                except Exception:
                    pass
                headings.append({"level": level, "text": text})
            else:
                paragraphs.append(text)

        if include_tables:
            for table in doc.tables:
                rows = []
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells]
                    if any(row_cells):
                        rows.append(row_cells)
                if rows:
                    tables_data.append(rows)

        preview = f"Read '{os.path.basename(abs_path)}': {len(headings)} heading(s), {len(paragraphs)} paragraph(s), {len(tables_data)} table(s)."
        return ToolExecutionResult(
            success=True,
            tool_id=self.id,
            data={
                "filename": os.path.basename(abs_path),
                "total_paragraphs": len(paragraphs),
                "total_headings": len(headings),
                "total_tables": len(tables_data),
                "headings": headings,
                "paragraphs": paragraphs[:50],  # preview top 50 paragraphs
                "tables": tables_data
            },
            output_preview=preview
        )


# ====================================================================
# 3. Excel Processor
# ====================================================================
class ExcelProcessorTool(BaseTool):
    id = "excel_processor"
    name = "Excel Processor"
    description = "Inspects and processes Excel (.xlsx, .xls) files. Reads sheets, displays columns, filters, sorts, calculates aggregations, and exports downloadable processed workbooks."
    category = "Documents & Files"
    icon = "Sheet"
    timeout_seconds = 30
    permission_required = "creator"

    input_schema = {
        "type": "object",
        "properties": {
            "file_path": {
                "type": "string",
                "description": "File name or path of the uploaded Excel file."
            },
            "sheet_name": {
                "type": "string",
                "description": "Optional name of the worksheet to inspect or operate on."
            },
            "operation": {
                "type": "string",
                "enum": ["info", "filter", "sort", "aggregate", "clean"],
                "description": "Operation: 'info' (view sheets and columns), 'filter', 'sort', 'aggregate', or 'clean'."
            },
            "column": {
                "type": "string",
                "description": "Target column name for filter, sort, or aggregate."
            },
            "filter_value": {
                "type": "string",
                "description": "Value to match when filtering."
            },
            "filter_operator": {
                "type": "string",
                "enum": ["==", "!=", ">", "<", ">=", "<=", "contains"],
                "description": "Operator for filtering rows."
            },
            "sort_ascending": {
                "type": "boolean",
                "description": "True for ascending, False for descending."
            },
            "agg_function": {
                "type": "string",
                "enum": ["sum", "mean", "min", "max", "count"],
                "description": "Aggregation function to apply on the column."
            },
            "group_by_column": {
                "type": "string",
                "description": "Optional column to group by before aggregating."
            }
        },
        "required": ["file_path"]
    }

    output_schema = {
        "type": "object",
        "properties": {
            "sheets": {"type": "array"},
            "active_sheet": {"type": "string"},
            "columns": {"type": "array"},
            "total_rows": {"type": "integer"},
            "preview": {"type": "array"},
            "download_url": {"type": "string"}
        }
    }

    async def execute(self, params: Dict[str, Any], context: Optional[Dict[str, Any]] = None) -> ToolExecutionResult:
        file_path = params.get("file_path")
        if not file_path:
            return ToolExecutionResult(success=False, tool_id=self.id, error="Missing required parameter 'file_path'.")

        try:
            abs_path = resolve_safe_file_path(file_path)
            with pd.ExcelFile(abs_path) as excel_file:
                sheet_names = list(excel_file.sheet_names)
                target_sheet = params.get("sheet_name") or sheet_names[0]
                if target_sheet not in sheet_names:
                    return ToolExecutionResult(
                        success=False,
                        tool_id=self.id,
                        error=f"Sheet '{target_sheet}' not found. Available sheets: {sheet_names}"
                    )
                df = pd.read_excel(excel_file, sheet_name=target_sheet)
        except Exception as e:
            return ToolExecutionResult(success=False, tool_id=self.id, error=f"Could not open Excel file: {str(e)}")

        operation = (params.get("operation") or "info").lower()
        col = params.get("column")

        # 1. Info operation
        if operation == "info":
            preview = df.head(10).to_dict(orient="records")
            return ToolExecutionResult(
                success=True,
                tool_id=self.id,
                data={
                    "sheets": sheet_names,
                    "sheet_names": sheet_names,
                    "active_sheet": target_sheet,
                    "columns": list(df.columns.astype(str)),
                    "total_rows": len(df),
                    "preview": preview
                },
                output_preview=f"Excel file '{os.path.basename(abs_path)}' ({target_sheet}): {len(df)} rows, {len(df.columns)} columns."
            )

        # 2. Filter operation
        if operation == "filter":
            if not col or col not in df.columns:
                return ToolExecutionResult(success=False, tool_id=self.id, error=f"Target column '{col}' not found in sheet. Available: {list(df.columns)}")
            
            val = params.get("filter_value", "")
            op = params.get("filter_operator", "==")

            if op == "contains":
                mask = df[col].astype(str).str.contains(str(val), case=False, na=False)
            elif op == "==":
                mask = df[col].astype(str) == str(val)
            elif op == "!=":
                mask = df[col].astype(str) != str(val)
            elif op in (">", "<", ">=", "<="):
                try:
                    num_val = float(val)
                    num_series = pd.to_numeric(df[col], errors="coerce")
                    if op == ">": mask = num_series > num_val
                    elif op == "<": mask = num_series < num_val
                    elif op == ">=": mask = num_series >= num_val
                    elif op == "<=": mask = num_series <= num_val
                except Exception:
                    mask = df[col].astype(str) == str(val)
            else:
                mask = df[col].astype(str) == str(val)

            filtered_df = df[mask]
            out_buf = io.BytesIO()
            filtered_df.to_excel(out_buf, index=False, engine="openpyxl")
            _, download_url = save_output_file(out_buf.getvalue(), prefix="filtered_excel", ext=".xlsx")

            return ToolExecutionResult(
                success=True,
                tool_id=self.id,
                data={
                    "original_rows": len(df),
                    "matched_rows": len(filtered_df),
                    "columns": list(filtered_df.columns.astype(str)),
                    "preview": filtered_df.head(10).to_dict(orient="records"),
                    "download_url": download_url
                },
                download_url=download_url,
                output_preview=f"Filtered {len(filtered_df)} row(s) matching '{col} {op} {val}' out of {len(df)} total."
            )

        # 3. Sort operation
        if operation == "sort":
            if not col or col not in df.columns:
                return ToolExecutionResult(success=False, tool_id=self.id, error=f"Sort column '{col}' not found. Available: {list(df.columns)}")
            asc = params.get("sort_ascending", True)
            sorted_df = df.sort_values(by=col, ascending=asc)

            out_buf = io.BytesIO()
            sorted_df.to_excel(out_buf, index=False, engine="openpyxl")
            _, download_url = save_output_file(out_buf.getvalue(), prefix="sorted_excel", ext=".xlsx")

            return ToolExecutionResult(
                success=True,
                tool_id=self.id,
                data={
                    "sorted_by": col,
                    "ascending": asc,
                    "total_rows": len(sorted_df),
                    "preview": sorted_df.head(10).to_dict(orient="records"),
                    "download_url": download_url
                },
                download_url=download_url,
                output_preview=f"Sorted {len(sorted_df)} row(s) by '{col}' ({'ascending' if asc else 'descending'})."
            )

        # 4. Aggregate operation
        if operation == "aggregate":
            if not col or col not in df.columns:
                return ToolExecutionResult(success=False, tool_id=self.id, error=f"Aggregation column '{col}' not found.")
            agg_func = (params.get("agg_function") or "sum").lower()
            group_col = params.get("group_by_column")

            numeric_series = pd.to_numeric(df[col], errors="coerce")
            if group_col and group_col in df.columns:
                df["_metric"] = numeric_series
                grouped = df.groupby(group_col)["_metric"].agg(agg_func).reset_index()
                grouped.rename(columns={"_metric": f"{agg_func}_{col}"}, inplace=True)
                
                out_buf = io.BytesIO()
                grouped.to_excel(out_buf, index=False, engine="openpyxl")
                _, download_url = save_output_file(out_buf.getvalue(), prefix="agg_excel", ext=".xlsx")

                return ToolExecutionResult(
                    success=True,
                    tool_id=self.id,
                    data={
                        "grouped_by": group_col,
                        "aggregation": agg_func,
                        "column": col,
                        "results": grouped.to_dict(orient="records"),
                        "download_url": download_url
                    },
                    download_url=download_url,
                    output_preview=f"Aggregated {agg_func.upper()} of '{col}' grouped by '{group_col}'."
                )
            else:
                calc_val = getattr(numeric_series, agg_func)()
                return ToolExecutionResult(
                    success=True,
                    tool_id=self.id,
                    data={
                        "column": col,
                        "aggregation": agg_func,
                        "result": float(calc_val) if pd.notnull(calc_val) else None
                    },
                    output_preview=f"{agg_func.upper()} of '{col}': {calc_val}"
                )

        # 5. Clean operation
        if operation == "clean":
            clean_df = df.dropna(how="all").drop_duplicates()
            out_buf = io.BytesIO()
            clean_df.to_excel(out_buf, index=False, engine="openpyxl")
            _, download_url = save_output_file(out_buf.getvalue(), prefix="cleaned_excel", ext=".xlsx")

            return ToolExecutionResult(
                success=True,
                tool_id=self.id,
                data={
                    "original_rows": len(df),
                    "cleaned_rows": len(clean_df),
                    "rows_removed": len(df) - len(clean_df),
                    "preview": clean_df.head(10).to_dict(orient="records"),
                    "download_url": download_url
                },
                download_url=download_url,
                output_preview=f"Cleaned workbook: removed {len(df) - len(clean_df)} duplicate/empty rows."
            )

        return ToolExecutionResult(success=False, tool_id=self.id, error=f"Unknown operation '{operation}'.")


# ====================================================================
# 4. CSV Processor
# ====================================================================
class CsvProcessorTool(BaseTool):
    id = "csv_processor"
    name = "CSV Processor"
    description = "Inspects and processes CSV files. Auto-detects delimiters and encodings, previews data, filters, sorts, transforms columns, and exports processed CSV files."
    category = "Documents & Files"
    icon = "FileSpreadsheet"
    timeout_seconds = 25
    permission_required = "creator"

    input_schema = {
        "type": "object",
        "properties": {
            "file_path": {
                "type": "string",
                "description": "File name or path of the uploaded CSV."
            },
            "operation": {
                "type": "string",
                "enum": ["info", "filter", "sort", "transform"],
                "description": "Operation to perform: 'info', 'filter', 'sort', or 'transform'."
            },
            "delimiter": {
                "type": "string",
                "description": "Optional delimiter override (e.g. ',', ';', '\\t', '|')."
            },
            "column": {
                "type": "string",
                "description": "Target column for filtering or sorting."
            },
            "filter_value": {
                "type": "string",
                "description": "Value to match when filtering."
            },
            "sort_ascending": {
                "type": "boolean",
                "description": "Sort order: True for ascending, False for descending."
            },
            "select_columns": {
                "type": "array",
                "items": {"type": "string"},
                "description": "Optional subset of column names to select and keep."
            }
        },
        "required": ["file_path"]
    }

    output_schema = {
        "type": "object",
        "properties": {
            "delimiter_detected": {"type": "string"},
            "encoding_detected": {"type": "string"},
            "columns": {"type": "array"},
            "total_rows": {"type": "integer"},
            "preview": {"type": "array"},
            "download_url": {"type": "string"}
        }
    }

    async def execute(self, params: Dict[str, Any], context: Optional[Dict[str, Any]] = None) -> ToolExecutionResult:
        file_path = params.get("file_path")
        csv_content = params.get("csv_content") or params.get("content")

        if not file_path and not csv_content:
            return ToolExecutionResult(success=False, tool_id=self.id, error="Missing required parameter 'file_path' or 'csv_content'.")

        delimiter = params.get("delimiter") or ","
        encoding = "utf-8"

        try:
            if file_path:
                abs_path = resolve_safe_file_path(file_path)
                # Detect encoding
                try:
                    with open(abs_path, "r", encoding="utf-8") as f:
                        sample = f.read(4096)
                except UnicodeDecodeError:
                    encoding = "latin-1"
                    with open(abs_path, "r", encoding="latin-1") as f:
                        sample = f.read(4096)

                if not params.get("delimiter"):
                    try:
                        sniffer = csv.Sniffer()
                        delimiter = sniffer.sniff(sample, delimiters=";,|\t,").delimiter
                    except Exception:
                        delimiter = ","

                df = pd.read_csv(abs_path, sep=delimiter, encoding=encoding, on_bad_lines="skip")
            else:
                df = pd.read_csv(io.StringIO(csv_content), sep=delimiter, on_bad_lines="skip")
        except Exception as e:
            return ToolExecutionResult(success=False, tool_id=self.id, error=f"Failed to parse CSV: {str(e)}")

        operation = (params.get("operation") or "info").lower()
        col = params.get("column") or params.get("filter_column")

        # 1. Info operation
        if operation == "info":
            return ToolExecutionResult(
                success=True,
                tool_id=self.id,
                data={
                    "delimiter_detected": delimiter,
                    "encoding_detected": encoding,
                    "columns": list(df.columns.astype(str)),
                    "total_rows": len(df),
                    "row_count": len(df),
                    "sample_data": df.head(10).to_dict(orient="records"),
                    "preview": df.head(10).to_dict(orient="records")
                },
                output_preview=f"CSV dataset: {len(df)} rows, {len(df.columns)} columns (delimiter: '{delimiter}')."
            )

        # 2. Filter operation
        if operation == "filter":
            val = params.get("filter_value", "")
            if not col or col not in df.columns:
                return ToolExecutionResult(success=False, tool_id=self.id, error=f"Filter column '{col}' not found. Available: {list(df.columns)}")
            
            filter_op = (params.get("filter_operator") or "contains").lower()
            if filter_op in ("equals", "==", "exact"):
                mask = df[col].astype(str) == str(val)
            else:
                mask = df[col].astype(str).str.contains(str(val), case=False, na=False)
            filtered_df = df[mask]
            
            csv_str = filtered_df.to_csv(index=False)
            _, download_url = save_output_file(csv_str.encode("utf-8"), prefix="filtered_csv", ext=".csv")

            return ToolExecutionResult(
                success=True,
                tool_id=self.id,
                data={
                    "original_rows": len(df),
                    "matched_rows": len(filtered_df),
                    "row_count": len(filtered_df),
                    "columns": list(filtered_df.columns.astype(str)),
                    "sample_data": filtered_df.head(10).to_dict(orient="records"),
                    "preview": filtered_df.head(10).to_dict(orient="records"),
                    "download_url": download_url
                },
                download_url=download_url,
                output_preview=f"Filtered {len(filtered_df)} of {len(df)} rows where '{col}' matches '{val}'."
            )

        # 3. Sort operation
        if operation == "sort":
            if not col or col not in df.columns:
                return ToolExecutionResult(success=False, tool_id=self.id, error=f"Sort column '{col}' not found.")
            asc = params.get("sort_ascending", True)
            sorted_df = df.sort_values(by=col, ascending=asc)

            csv_str = sorted_df.to_csv(index=False)
            _, download_url = save_output_file(csv_str.encode("utf-8"), prefix="sorted_csv", ext=".csv")

            return ToolExecutionResult(
                success=True,
                tool_id=self.id,
                data={
                    "sorted_by": col,
                    "ascending": asc,
                    "total_rows": len(sorted_df),
                    "preview": sorted_df.head(10).to_dict(orient="records"),
                    "download_url": download_url
                },
                download_url=download_url,
                output_preview=f"Sorted {len(sorted_df)} row(s) by '{col}'."
            )

        # 4. Transform operation
        if operation == "transform":
            select_cols = params.get("select_columns")
            out_df = df
            if select_cols and isinstance(select_cols, list):
                valid_cols = [c for c in select_cols if c in df.columns]
                if valid_cols:
                    out_df = df[valid_cols]

            csv_str = out_df.to_csv(index=False)
            _, download_url = save_output_file(csv_str.encode("utf-8"), prefix="transformed_csv", ext=".csv")

            return ToolExecutionResult(
                success=True,
                tool_id=self.id,
                data={
                    "columns": list(out_df.columns.astype(str)),
                    "total_rows": len(out_df),
                    "preview": out_df.head(10).to_dict(orient="records"),
                    "download_url": download_url
                },
                download_url=download_url,
                output_preview=f"Transformed CSV with {len(out_df.columns)} columns and {len(out_df)} rows."
            )

        return ToolExecutionResult(success=False, tool_id=self.id, error=f"Unknown operation '{operation}'.")


# ====================================================================
# 5. JSON Parser
# ====================================================================
class JsonParserTool(BaseTool):
    id = "json_parser"
    name = "JSON Parser"
    description = "Parses, validates, formats, and queries JSON structures. Extracts nested keys using dot notation and filters arrays with detailed syntax error reporting."
    category = "Documents & Files"
    icon = "FileJson"
    timeout_seconds = 15
    permission_required = "creator"

    input_schema = {
        "type": "object",
        "properties": {
            "json_data": {
                "type": "string",
                "description": "Raw JSON string or file path containing the JSON to parse."
            },
            "operation": {
                "type": "string",
                "enum": ["validate", "extract", "filter", "format"],
                "description": "Operation: 'validate', 'extract', 'filter', or 'format'."
            },
            "path": {
                "type": "string",
                "description": "Dot-notation field path to extract (e.g. 'users.0.name' or 'data.items')."
            },
            "filter_key": {
                "type": "string",
                "description": "Key to evaluate when filtering an array of objects."
            },
            "filter_value": {
                "type": "string",
                "description": "Expected value to match in filter array."
            }
        },
        "required": ["json_data"]
    }

    output_schema = {
        "type": "object",
        "properties": {
            "is_valid": {"type": "boolean"},
            "parsed_data": {"type": "object"},
            "result": {"type": "any"},
            "error_details": {"type": "object"}
        }
    }

    async def execute(self, params: Dict[str, Any], context: Optional[Dict[str, Any]] = None) -> ToolExecutionResult:
        raw_input = params.get("json_data") or params.get("json_content") or params.get("content")
        if not raw_input:
            return ToolExecutionResult(success=False, tool_id=self.id, error="Missing required parameter 'json_data' or 'json_content'.")

        json_text = raw_input.strip()

        # If it looks like a filename, try reading from disk
        if not (json_text.startswith("{") or json_text.startswith("[")) and (json_text.endswith(".json") or len(json_text) < 255):
            try:
                abs_path = resolve_safe_file_path(json_text)
                with open(abs_path, "r", encoding="utf-8") as f:
                    json_text = f.read()
            except Exception:
                pass  # treat as raw string

        # Validate syntax
        try:
            parsed = json.loads(json_text)
        except json.JSONDecodeError as decode_err:
            return ToolExecutionResult(
                success=False,
                tool_id=self.id,
                error=f"JSON Syntax Error at line {decode_err.lineno}, column {decode_err.colno}: {decode_err.msg}",
                data={
                    "valid": False,
                    "is_valid": False,
                    "error_line": decode_err.lineno,
                    "error_column": decode_err.colno,
                    "message": decode_err.msg
                }
            )

        operation = (params.get("operation") or ("extract" if (params.get("extract_path") or params.get("path")) else "validate")).lower()

        # 1. Validate / Format
        if operation in ("validate", "format"):
            formatted = json.dumps(parsed, indent=2)
            return ToolExecutionResult(
                success=True,
                tool_id=self.id,
                data={
                    "valid": True,
                    "is_valid": True,
                    "type": "object" if isinstance(parsed, dict) else "array" if isinstance(parsed, list) else "scalar",
                    "item_count": len(parsed) if isinstance(parsed, (dict, list)) else 1,
                    "formatted": formatted
                },
                output_preview=f"JSON is valid. Root type: {type(parsed).__name__} ({len(parsed) if isinstance(parsed, (dict, list)) else 1} items)."
            )

        # 2. Extract nested field (dot notation)
        if operation == "extract":
            path = params.get("extract_path") or params.get("path", "")
            if not path:
                return ToolExecutionResult(success=True, tool_id=self.id, data={"valid": True, "extracted_value": parsed}, output_preview="No path provided; returning full root.")

            current = parsed
            tokens = path.replace("[", ".").replace("]", "").split(".")
            for tok in tokens:
                if not tok:
                    continue
                if isinstance(current, dict):
                    if tok in current:
                        current = current[tok]
                    else:
                        return ToolExecutionResult(success=False, tool_id=self.id, error=f"Key '{tok}' not found in object at path '{path}'.")
                elif isinstance(current, list):
                    try:
                        idx = int(tok)
                        current = current[idx]
                    except (ValueError, IndexError):
                        return ToolExecutionResult(success=False, tool_id=self.id, error=f"Invalid array index '{tok}' in path '{path}'.")
                else:
                    return ToolExecutionResult(success=False, tool_id=self.id, error=f"Cannot traverse into scalar value at '{tok}'.")

            return ToolExecutionResult(
                success=True,
                tool_id=self.id,
                data={"valid": True, "path": path, "extracted_value": current},
                output_preview=f"Extracted '{path}': {json.dumps(current)[:150]}"
            )

        # 3. Filter array of objects
        if operation == "filter":
            if not isinstance(parsed, list):
                return ToolExecutionResult(success=False, tool_id=self.id, error="Filter operation requires a root JSON array of objects.")
            
            f_key = params.get("filter_key")
            f_val = params.get("filter_value")
            if not f_key:
                return ToolExecutionResult(success=False, tool_id=self.id, error="Missing parameter 'filter_key' for array filter.")

            filtered = [
                item for item in parsed
                if isinstance(item, dict) and str(item.get(f_key, "")).lower() == str(f_val).lower()
            ]

            return ToolExecutionResult(
                success=True,
                tool_id=self.id,
                data={
                    "total_items": len(parsed),
                    "matched_items": len(filtered),
                    "filtered_array": filtered
                },
                output_preview=f"Filtered {len(filtered)} of {len(parsed)} items where '{f_key}' == '{f_val}'."
            )

        return ToolExecutionResult(success=False, tool_id=self.id, error=f"Unknown operation '{operation}'.")
